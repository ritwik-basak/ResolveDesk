import csv
import io
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path

import psycopg2
from dotenv import load_dotenv
from llama_index.core import Document, StorageContext, VectorStoreIndex
from llama_index.core.ingestion import IngestionPipeline
from llama_index.core.node_parser import SentenceSplitter
from llama_index.core.schema import TextNode
from llama_index.vector_stores.pinecone import PineconeVectorStore
from pinecone import Pinecone, ServerlessSpec
from supabase import create_client

from backend.models import embed_model as _embed_model
from backend.models import EMBED_DIM

load_dotenv()


# =============================================================================
# CONFIGURATION
# =============================================================================

KNOWLEDGE_BASE_DIR = Path(__file__).parent / "knowledge_base"

PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
PINECONE_INDEX   = os.getenv("PINECONE_INDEX_NAME", "resolvedesk-knowledge")
PINECONE_REGION  = os.getenv("PINECONE_ENVIRONMENT", "us-east-1")
CHUNK_SIZE    = 500
CHUNK_OVERLAP = 100

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY")
DB_URL       = os.getenv("SUPABASE_DB_URL")


# =============================================================================
# SINGLETONS — loaded once, reused across all ingest calls
# =============================================================================

_splitter = SentenceSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
_supabase    = create_client(SUPABASE_URL, SUPABASE_KEY)


def _get_pg():
    return psycopg2.connect(DB_URL)


def _get_pinecone_index():
    pc       = Pinecone(api_key=PINECONE_API_KEY)
    existing = [i.name for i in pc.list_indexes()]
    if PINECONE_INDEX not in existing:
        pc.create_index(
            name      = PINECONE_INDEX,
            dimension = EMBED_DIM,
            metric    = "cosine",
            spec      = ServerlessSpec(cloud="aws", region=PINECONE_REGION),
        )
    return pc.Index(PINECONE_INDEX)


# =============================================================================
# TABLE SETUP
# Called once at startup in main.py. Safe to call every time (IF NOT EXISTS).
# =============================================================================

_CREATE_DOCUMENTS_TABLE = """
CREATE TABLE IF NOT EXISTS documents (
    doc_id      TEXT PRIMARY KEY,
    filename    TEXT NOT NULL UNIQUE,
    upload_date TIMESTAMPTZ NOT NULL DEFAULT NOW(),
    chunk_count INTEGER NOT NULL,
    gcs_path    TEXT NOT NULL,
    file_type   TEXT NOT NULL
);
"""

_CREATE_CHUNKS_TABLE = """
CREATE TABLE IF NOT EXISTS chunks (
    chunk_id    TEXT PRIMARY KEY,
    text        TEXT NOT NULL,
    source      TEXT NOT NULL,
    category    TEXT NOT NULL,
    chunk_index INTEGER NOT NULL
);
"""


def create_document_tables() -> None:
    """Create documents + chunks tables if they don't exist. Called at startup."""
    conn = _get_pg()
    cur  = conn.cursor()
    cur.execute(_CREATE_DOCUMENTS_TABLE)
    cur.execute(_CREATE_CHUNKS_TABLE)
    conn.commit()
    cur.close()
    conn.close()
    print("[Ingest] Tables ready (documents, chunks).")


# =============================================================================
# HELPERS
# =============================================================================

def _get_category(source: str) -> str:
    if "faq" in source:
        return "faq"
    if "policy" in source:
        return "policy"
    return "support"


# =============================================================================
# FILE LOADERS
# Each loader returns LlamaIndex Document objects from raw file bytes.
#
# PDF / DOCX / PPTX / TXT / MD → one Document per file (SentenceSplitter chunks it)
# CSV                           → one Document per row  (already pre-chunked)
# =============================================================================

def _load_pdf(file_bytes: bytes, filename: str) -> list[Document]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    text   = "\n".join(p.extract_text() or "" for p in reader.pages)
    return [Document(text=text, metadata={"filename": filename})]


def _load_docx(file_bytes: bytes, filename: str) -> list[Document]:
    import docx
    doc  = docx.Document(io.BytesIO(file_bytes))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [Document(text=text, metadata={"filename": filename})]


def _load_pptx(file_bytes: bytes, filename: str) -> list[Document]:
    from pptx import Presentation
    prs   = Presentation(io.BytesIO(file_bytes))
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if para.text.strip():
                        lines.append(para.text.strip())
    return [Document(text="\n".join(lines), metadata={"filename": filename})]


def _load_txt_md(file_bytes: bytes, filename: str) -> list[Document]:
    text = file_bytes.decode("utf-8", errors="replace")
    return [Document(text=text, metadata={"filename": filename})]


def _load_csv(file_bytes: bytes, filename: str) -> list[Document]:
    """One Document per CSV row. Headers prepended: "col1: val1 | col2: val2"."""
    text   = file_bytes.decode("utf-8", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    docs   = []
    for row in reader:
        chunk_text = " | ".join(f"{k}: {v}" for k, v in row.items() if v)
        if chunk_text.strip():
            docs.append(Document(text=chunk_text, metadata={"filename": filename}))
    return docs


_LOADERS = {
    ".pdf":  _load_pdf,
    ".docx": _load_docx,
    ".pptx": _load_pptx,
    ".txt":  _load_txt_md,
    ".md":   _load_txt_md,
    ".csv":  _load_csv,
}

SUPPORTED_EXTENSIONS = set(_LOADERS.keys())


def _load_file(file_bytes: bytes, filename: str) -> list[Document]:
    ext = Path(filename).suffix.lower()
    if ext not in _LOADERS:
        raise ValueError(f"Unsupported file type: {ext}")
    return _LOADERS[ext](file_bytes, filename)


# =============================================================================
# CORE INGEST FUNCTION
# =============================================================================

def ingest_file(file_bytes: bytes, filename: str, gcs_path: str) -> int:
    """
    Run the full LlamaIndex ingestion pipeline for one file.

    PIPELINE (6 steps):
      1. Load   — file bytes → LlamaIndex Documents via type-specific loader
      2. Split  — SentenceSplitter (500/100) or row-per-node for CSV
      3. Tag    — assign source, category, chunk_index metadata to each node
      4. Embed  — BGE-M3 + upsert to Pinecone via LlamaIndex IngestionPipeline
      5. Chunks — save text + metadata to Supabase chunks table (for BM25)
      6. Record — save document metadata to Supabase documents table

    Returns:
        Number of chunks created.
    """
    source    = Path(filename).stem
    category  = _get_category(source)
    file_type = Path(filename).suffix.lower().lstrip(".")

    # ── Step 1: Load ─────────────────────────────────────────────────────
    raw_docs = _load_file(file_bytes, filename)

    # ── Step 2: Split ────────────────────────────────────────────────────
    # CSV rows are already individual chunks — wrap as TextNodes directly.
    # All other types: SentenceSplitter breaks Documents into overlapping chunks.
    if file_type == "csv":
        nodes = [
            TextNode(text=doc.text, metadata=dict(doc.metadata))
            for doc in raw_docs
        ]
    else:
        nodes = _splitter.get_nodes_from_documents(raw_docs)

    # ── Step 3: Assign metadata ──────────────────────────────────────────
    for i, node in enumerate(nodes):
        node.id_                     = f"{source}_chunk_{i}"
        node.metadata["source"]      = source
        node.metadata["category"]    = category
        node.metadata["chunk_index"] = i

    # ── Step 4: Embed + upsert to Pinecone (LlamaIndex IngestionPipeline) ─
    pinecone_idx = _get_pinecone_index()
    vector_store = PineconeVectorStore(pinecone_index=pinecone_idx)

    pipeline = IngestionPipeline(
        transformations=[_embed_model],
        vector_store=vector_store,
    )
    pipeline.run(nodes=nodes, show_progress=True)

    # ── Step 5: Save chunks to Supabase (rebuilds BM25 in retriever) ─────
    chunk_rows = [
        {
            "chunk_id":    node.id_,
            "text":        node.text,
            "source":      source,
            "category":    category,
            "chunk_index": node.metadata["chunk_index"],
        }
        for node in nodes
    ]
    _supabase.table("chunks").upsert(chunk_rows).execute()

    # ── Step 6: Save document metadata ───────────────────────────────────
    _supabase.table("documents").upsert([{
        "doc_id":      str(uuid.uuid4()),
        "filename":    filename,
        "upload_date": datetime.now(timezone.utc).isoformat(),
        "chunk_count": len(nodes),
        "gcs_path":    gcs_path,
        "file_type":   file_type,
    }], on_conflict="filename").execute()

    return len(nodes)


# =============================================================================
# DELETE FUNCTION
# Removes all data for a document from Pinecone + Supabase.
# GCS deletion is handled separately by the router (gcs.py).
# =============================================================================

def delete_file_data(filename: str) -> None:
    """
    Delete all traces of a document:
      - Pinecone vectors  (filter: source == filename stem)
      - Supabase chunks   (DELETE WHERE source == filename stem)
      - Supabase documents (DELETE WHERE filename == filename)
    """
    source = Path(filename).stem

    pc_idx = _get_pinecone_index()
    pc_idx.delete(filter={"source": {"$eq": source}})

    _supabase.table("chunks").delete().eq("source", source).execute()
    _supabase.table("documents").delete().eq("filename", filename).execute()


# =============================================================================
# INITIAL SETUP SCRIPT
# Ingest all .md files from knowledge_base/ into Pinecone + Supabase.
#
# Run this once after deployment (or after any knowledge base update):
#   python backend/rag/ingest.py
# =============================================================================

def main():
    print("=" * 60)
    print("  ResolveDesk RAG Ingest (LlamaIndex)")
    print("=" * 60)

    create_document_tables()

    md_files = sorted(KNOWLEDGE_BASE_DIR.glob("*.md"))
    print(f"\nFound {len(md_files)} documents in knowledge_base/\n")

    total_chunks = 0
    for md_file in md_files:
        print(f"  Ingesting {md_file.name}...")
        n = ingest_file(
            file_bytes=md_file.read_bytes(),
            filename=md_file.name,
            gcs_path=f"knowledge_base/{md_file.name}",
        )
        print(f"    -> {n} chunks")
        total_chunks += n

    print(f"\n{'=' * 60}")
    print(f"  Done. {total_chunks} total chunks ingested.")
    print(f"{'=' * 60}")


if __name__ == "__main__":
    main()
